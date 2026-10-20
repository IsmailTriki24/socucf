import io
import json

import openpyxl
import pytest

from accounts.permissions import CLIENT_PERMISSIONS, can
from audit.models import JournalAudit
from catalog.models import Domaine, SourceLog, Template, render_parametres
from clients.models import Client, ClientSource, DetectionClient
from tests.conftest import api, make_client


def get_det(client, code="DET-001"):
    return DetectionClient.objects.get(client=client, code=code)


def tech_ids(matrix):
    return {t["id"]: t for tac in matrix["tactiques"] for t in tac["techniques"]}


# ---------------------------------------------------------------- création du client
@pytest.mark.django_db
def test_client_creation_copies_all_domains_by_default(world):
    r = api(world.admin).post("/api/clients", {"nom": "Tout coché", "siems": ["wazuh", "fortisiem"]}, format="json")
    assert r.status_code == 201
    body = r.json()
    assert body["nb_detections"] == 82 and len(body["domaines"]) == Domaine.objects.count()
    assert len(body["bande"]) == 82 and body["par_statut"]["brouillon"] == 82
    c = Client.objects.get(pk=body["id"])
    det = get_det(c)
    assert det.syntaxes.count() == 1 and det.syntaxes.first().siem.code == "wazuh"  # règle copiée pour la plateforme du client
    assert det.detection_sources.count() == 4 and det.techniques.count() == 2


@pytest.mark.django_db
def test_client_creation_requires_name_platform_and_domain(world):
    c = api(world.admin)
    assert c.post("/api/clients", {"nom": " ", "siems": ["wazuh"]}, format="json").status_code == 400
    assert c.post("/api/clients", {"nom": "X", "siems": []}, format="json").status_code == 400
    assert c.post("/api/clients", {"nom": "X", "siems": ["wazuh"], "domaines": []}, format="json").status_code == 400
    assert c.post("/api/clients", {"nom": "client a", "siems": ["wazuh"]}, format="json").status_code == 409


@pytest.mark.django_db
def test_client_with_chosen_domains_only(world):
    d = Domaine.objects.get(nom="Identité & AD")
    r = api(world.admin).post("/api/clients", {"nom": "Un domaine", "siems": ["wazuh"], "domaines": [d.pk]}, format="json").json()
    n = Template.objects.filter(domaine=d).count()
    assert r["nb_detections"] == n and n < 82


@pytest.mark.django_db
def test_adding_a_domain_adds_its_templates_and_a_platform_adds_rules(world):
    a = api(world.adminclient)
    dom = Domaine.objects.exclude(pk__in=world.a.domaines.all()).first()
    keep = list(world.a.domaines.values_list("pk", flat=True))
    before = DetectionClient.objects.filter(client=world.a).count()
    r = a.patch(f"/api/clients/{world.a.pk}", {"domaines": keep + [dom.pk]}, format="json")
    assert r.status_code == 200 and r.json()["detections_creees"] == Template.objects.filter(domaine=dom).count()
    assert DetectionClient.objects.filter(client=world.a).count() == before + r.json()["detections_creees"]
    # on garde les détections existantes quand un domaine est retiré
    a.patch(f"/api/clients/{world.a.pk}", {"domaines": keep}, format="json")
    assert DetectionClient.objects.filter(client=world.a).count() == before + r.json()["detections_creees"]
    # nouvelle plateforme : les règles du template sont complétées
    assert get_det(world.a).syntaxes.filter(siem__code="splunk").count() == 0
    r = a.patch(f"/api/clients/{world.a.pk}", {"siems": ["wazuh", "sentinel", "splunk"]}, format="json")
    assert r.json()["regles_ajoutees"] >= 1 and get_det(world.a).syntaxes.filter(siem__code="splunk").count() == 1


@pytest.mark.django_db
def test_only_admin_deletes_client_with_name_confirmation(world):
    assert api(world.adminclient).delete(f"/api/clients/{world.a.pk}", {"confirmation": world.a.nom}, format="json").status_code == 403
    assert api(world.admin).delete(f"/api/clients/{world.a.pk}", {"confirmation": "autre"}, format="json").status_code == 400
    assert api(world.admin).delete(f"/api/clients/{world.a.pk}", {"confirmation": world.a.nom}, format="json").status_code == 204
    assert not Client.objects.filter(pk=world.a.pk).exists() and Client.objects.filter(pk=world.b.pk).exists()


# ---------------------------------------------------------------- détections : tout est modifiable
@pytest.mark.django_db
def test_status_is_free_and_deploy_sets_date(world):
    d = get_det(world.a)
    c = api(world.contrib)
    for statut in ("deployee", "a_revoir", "validee", "en_test", "brouillon", "deployee"):  # aucune règle de passage
        r = c.patch(f"/api/detections/{d.pk}", {"statut": statut}, format="json")
        assert r.status_code == 200 and r.json()["statut"] == statut
    assert r.json()["date_deploiement"]
    assert c.patch(f"/api/detections/{d.pk}", {"statut": "inconnu"}, format="json").status_code == 400
    assert api(world.lecteur).patch(f"/api/detections/{d.pk}", {"statut": "validee"}, format="json").status_code == 403


@pytest.mark.django_db
def test_exclusion_recommendation_impact_and_all_content_editable(world):
    d = get_det(world.a)
    body = {"exclusions": "comptes de service", "recommandation": "verrouiller après 5 échecs", "impact": "compte compromis",
            "cas_usage": "Mon cas d'usage", "gravite": "Critique", "type_detection": "Seuil", "logique": "SI x ALORS y",
            "logs_a_collecter": "4625", "parametres": {"seuil": 12}, "techniques": ["T1110.001"],
            "sources": [{"source_id": SourceLog.objects.first().pk, "obligatoire": True}]}
    r = api(world.valideur).patch(f"/api/detections/{d.pk}", body, format="json")
    assert r.status_code == 200
    j = r.json()
    assert (j["exclusions"], j["recommandation"], j["impact"], j["gravite"]) == ("comptes de service", "verrouiller après 5 échecs", "compte compromis", "Critique")
    assert [t["identifiant"] for t in j["techniques"]] == ["T1110.001"] and len(j["sources"]) == 1 and j["sources"][0]["obligatoire"]
    assert j["parametres_effectifs"]["seuil"] == 12
    # le template de la bibliothèque n'est pas touché
    assert Template.objects.get(code="DET-001").cas_usage != "Mon cas d'usage"
    bad = api(world.valideur).patch(f"/api/detections/{d.pk}", {"techniques": ["T9999"]}, format="json")
    assert bad.status_code == 400 and "T9999" in bad.json()["detail"]


@pytest.mark.django_db
def test_rules_hunting_add_and_delete(world):
    d = get_det(world.a)
    c = api(world.contrib)
    r = c.put(f"/api/detections/{d.pk}/syntaxes/sentinel", {"contenu": "SecurityEvent | take 1 // {{seuil}}"}, format="json")
    assert r.status_code == 200
    s = next(x for x in r.json()["syntaxes"] if x["siem"] == "sentinel")
    assert s["presente"] and "{{seuil}}" in s["contenu"] and "8" in s["rendu"]
    assert c.put(f"/api/detections/{d.pk}/syntaxes/splunk", {"contenu": "x"}, format="json").status_code == 400  # plateforme non utilisée
    assert next(x for x in c.delete(f"/api/detections/{d.pk}/syntaxes/sentinel").json()["syntaxes"] if x["siem"] == "sentinel")["presente"] is False
    r = c.post(f"/api/detections/{d.pk}/hunting", {"siem": "wazuh", "titre": "Chasse", "requete": "q"}, format="json")
    assert r.status_code == 201
    hid = r.json()["hunting"][-1]["id"]
    assert c.patch(f"/api/detections/{d.pk}/hunting/{hid}", {"titre": "Renommée"}, format="json").json()["hunting"][-1]["titre"] == "Renommée"
    assert c.delete(f"/api/detections/{d.pk}/hunting/{hid}").json()["hunting"] == [h for h in r.json()["hunting"] if h["id"] != hid]


@pytest.mark.django_db
def test_add_from_library_custom_and_delete(world):
    c = api(world.contrib)
    d = get_det(world.a)
    assert c.delete(f"/api/detections/{d.pk}").status_code == 204
    lib = c.get(f"/api/clients/{world.a.pk}/bibliotheque").json()
    assert any(t["code"] == "DET-001" for t in lib)
    assert c.post(f"/api/clients/{world.a.pk}/detections", {"codes": ["DET-001"]}, format="json").status_code == 201
    assert c.post(f"/api/clients/{world.a.pk}/detections", {"codes": ["DET-001"]}, format="json").status_code == 400  # déjà là
    dom = Domaine.objects.first()
    r = c.post(f"/api/clients/{world.a.pk}/detections", {"libre": {"domaine_id": dom.pk, "cas_usage": "Spécifique client", "gravite": "Haute"}}, format="json")
    assert r.status_code == 201
    libre = DetectionClient.objects.get(pk=r.json()["ids"][0])
    assert libre.code == "CLI-001" and libre.template is None and libre.gravite == "Haute"
    assert api(world.lecteur).post(f"/api/clients/{world.a.pk}/detections", {"codes": ["DET-002"]}, format="json").status_code == 403


@pytest.mark.django_db
def test_bulk_actions(world):
    c = api(world.contrib)
    ids = list(DetectionClient.objects.filter(client=world.a).values_list("pk", flat=True)[:3])
    assert c.post(f"/api/clients/{world.a.pk}/detections/groupe", {"ids": ids, "action": "statut", "valeur": "en_test"}, format="json").json()["modifiees"] == 3
    assert DetectionClient.objects.filter(pk__in=ids, statut="en_test").count() == 3
    c.post(f"/api/clients/{world.a.pk}/detections/groupe", {"ids": ids, "action": "journalisation", "valeur": "ok"}, format="json")
    assert DetectionClient.objects.filter(pk__in=ids, journalisation="ok").count() == 3
    c.post(f"/api/clients/{world.a.pk}/detections/groupe", {"ids": ids, "action": "supprimer"}, format="json")
    assert not DetectionClient.objects.filter(pk__in=ids).exists()
    # on ne touche pas aux détections d'un autre client
    other = list(DetectionClient.objects.filter(client=world.b).values_list("pk", flat=True)[:2])
    assert c.post(f"/api/clients/{world.a.pk}/detections/groupe", {"ids": other, "action": "supprimer"}, format="json").status_code == 400
    assert DetectionClient.objects.filter(pk__in=other).count() == 2


@pytest.mark.django_db
def test_journalisation_and_adding_a_log_source(world):
    # Client A collecte des sources utiles à DET-001 ; un client sans source collectée démarre en « Non »
    assert get_det(world.a).journalisation == "ok"
    pauvre = make_client("Sans logs", collectees=[])
    assert get_det(pauvre).journalisation == "non"
    c = api(world.admin)
    d = get_det(pauvre)
    src = d.detection_sources.first().source
    dash = c.get(f"/api/clients/{pauvre.pk}/dashboard").json()
    assert any(s["source_id"] == src.pk and s["nb_detections"] >= 1 for s in dash["sources_a_ajouter"])
    r = c.post(f"/api/clients/{pauvre.pk}/sources/{src.pk}/ajouter", {}, format="json")
    assert r.status_code == 200 and r.json()["detections_passees_ok"] >= 1
    assert get_det(pauvre).journalisation == "ok"
    assert ClientSource.objects.get(client=pauvre, source=src).statut == "collectee"
    # la journalisation reste modifiable à la main
    assert c.patch(f"/api/detections/{d.pk}", {"journalisation": "non"}, format="json").json()["journalisation"] == "non"
    assert c.patch(f"/api/detections/{d.pk}", {"journalisation": "peut-être"}, format="json").status_code == 400


@pytest.mark.django_db
def test_reset_from_template_and_template_deletion_keeps_client_detection(world):
    d = get_det(world.a)
    api(world.contrib).patch(f"/api/detections/{d.pk}", {"cas_usage": "Modifié", "impact": "x"}, format="json")
    api(world.contrib).put(f"/api/detections/{d.pk}/syntaxes/wazuh", {"contenu": "règle ajoutée par le client"}, format="json")
    api(world.contrib).post(f"/api/detections/{d.pk}/hunting", {"siem": "wazuh", "titre": "chasse client", "requete": "q"}, format="json")
    r = api(world.contrib).post(f"/api/detections/{d.pk}/reinitialiser", {}, format="json")
    assert r.status_code == 200 and r.json()["cas_usage"] == Template.objects.get(code="DET-001").cas_usage
    assert "règle ajoutée par le client" not in next(s["contenu"] for s in r.json()["syntaxes"] if s["siem"] == "wazuh")  # retour au template
    assert all(h["titre"] != "chasse client" for h in r.json()["hunting"])
    assert api(world.catalogue).delete("/api/templates/DET-001").status_code == 204
    d.refresh_from_db()
    assert d.template is None and d.cas_usage  # la détection du client survit
    assert api(world.contrib).post(f"/api/detections/{d.pk}/reinitialiser", {}, format="json").status_code == 400


# ---------------------------------------------------------------- couverture, exports, audit
@pytest.mark.django_db
def test_coverage_counts_only_chosen_status(world):
    d1 = get_det(world.a, "DET-001")
    api(world.contrib).patch(f"/api/detections/{d1.pk}", {"statut": "deployee"}, format="json")
    c = api(world.lecteur)
    default = tech_ids(c.get(f"/api/clients/{world.a.pk}/couverture").json())
    assert default["T1110"]["sous_techniques"][0]["detections"] == ["DET-001"]
    assert default["T1078"]["detections"] == ["DET-001"]  # les brouillons ne comptent pas par défaut
    both = tech_ids(c.get(f"/api/clients/{world.a.pk}/couverture?statuts=deployee,brouillon").json())
    assert len(both["T1078"]["detections"]) > 1
    r = c.get(f"/api/clients/{world.a.pk}/couverture").json()
    uncovered = {u["identifiant"]: u for u in r["non_couvertes"]}
    assert "T1078" not in uncovered and any(u["templates"] for u in uncovered.values())


@pytest.mark.django_db
def test_navigator_layer(world):
    api(world.contrib).patch(f"/api/detections/{get_det(world.a).pk}", {"statut": "deployee"}, format="json")
    r = api(world.lecteur).get(f"/api/clients/{world.a.pk}/export.navigator.json")
    layer = json.loads(r.content)
    assert layer["domain"] == "enterprise-attack" and layer["versions"]["layer"]
    by_id = {t["techniqueID"]: t for t in layer["techniques"]}
    assert by_id["T1078"]["score"] == 1 and by_id["T1078"]["comment"] == "DET-001"


@pytest.mark.django_db
def test_excel_export_has_catalogue_columns_plus_client_columns(world):
    det = get_det(world.a)
    api(world.contrib).patch(f"/api/detections/{det.pk}", {"parametres": {"seuil": 12}, "impact": "Compte compromis"}, format="json")
    r = api(world.lecteur).get(f"/api/clients/{world.a.pk}/export.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Detections"]
    header = [c.value for c in ws[1]]
    assert header[:9] == ["ID", "Domaine", "Cas d'usage (attaque détectée)", "Type", "Logs à collecter", "Source des logs", "Logique de la règle", "Gravité", "MITRE ATT&CK"]
    assert header[9:] == ["Statut", "Journalisation", "SIEM / XDR", "Règle", "Requêtes de hunting", "Exclusions", "Recommandation", "Impact", "Date de déploiement"]
    row = next([c.value for c in r_] for r_ in ws.iter_rows(min_row=2) if r_[0].value == "DET-001")
    assert row[9] == "Brouillon" and row[10] == "OK"
    assert 'frequency="12"' in row[12] and "{{seuil}}" not in row[12]
    assert "[Wazuh]" in row[12] and "[Microsoft Sentinel / Defender XDR]" in row[12] and row[16] == "Compte compromis"


@pytest.mark.django_db
def test_excel_and_csv_formula_injection_neutralised(world):
    from catalog.services.exports import put
    from clients.services.exports import _csv_safe

    ws = openpyxl.Workbook().active
    assert put(ws, 1, 1, '=HYPERLINK("http://evil","x")').data_type == "s"
    assert _csv_safe("=1+1").startswith("'") and _csv_safe("@SUM(A1)").startswith("'") and _csv_safe("-2+3").startswith("'")
    det = get_det(world.a)
    api(world.contrib).put(f"/api/detections/{det.pk}/syntaxes/wazuh", {"contenu": '=HYPERLINK("http://evil","x")'}, format="json")
    import csv

    text = api(world.lecteur).get(f"/api/clients/{world.a.pk}/export.csv").content.decode("utf-8").lstrip("﻿")
    for row in csv.reader(io.StringIO(text), delimiter=";"):
        assert all(not c.startswith(("=", "+", "-", "@")) for c in row)


@pytest.mark.django_db
def test_changes_are_audited(world):
    d = get_det(world.a)
    api(world.contrib).patch(f"/api/detections/{d.pk}", {"statut": "en_test"}, format="json")
    e = JournalAudit.objects.filter(client=world.a, objet_type="clients.DetectionClient", objet_id=str(d.pk)).first()
    assert e.action == "modification" and e.apres["statut"] == "en_test" and e.utilisateur == world.contrib
    api(world.contrib).delete(f"/api/detections/{d.pk}")
    assert JournalAudit.objects.filter(client=world.a, objet_id=str(d.pk), action="suppression").exists()


@pytest.mark.django_db
def test_render_parametres():
    out, missing = render_parametres("a {{seuil}} b {{ fenetre }} c {{inconnu}}", {"seuil": 5, "fenetre": "10m"})
    assert out == "a 5 b 10m c {{inconnu}}" and missing == ["inconnu"]


@pytest.mark.django_db
def test_permission_matrix(world):
    expect = {
        "client.view": {"admin_p": 1, "admin_a": 1, "valideur_a": 1, "contrib_a": 1, "lecteur_a": 1, "resp_cat": 0},
        "detection.edit": {"admin_p": 1, "admin_a": 1, "valideur_a": 1, "contrib_a": 1, "lecteur_a": 0, "resp_cat": 0},
        "habilitation.manage": {"admin_p": 1, "admin_a": 1, "valideur_a": 0, "contrib_a": 0, "lecteur_a": 0, "resp_cat": 0},
        "export": {"admin_p": 1, "admin_a": 1, "valideur_a": 1, "contrib_a": 1, "lecteur_a": 1, "resp_cat": 0},
        "audit.view": {"admin_p": 1, "admin_a": 1, "valideur_a": 1, "contrib_a": 0, "lecteur_a": 0, "resp_cat": 0},
    }
    users = {"admin_p": world.admin, "admin_a": world.adminclient, "valideur_a": world.valideur, "contrib_a": world.contrib,
             "lecteur_a": world.lecteur, "resp_cat": world.catalogue}
    assert set(expect) <= set(CLIENT_PERMISSIONS)
    for perm, row in expect.items():
        for name, allowed in row.items():
            assert can(users[name], world.a, perm) == bool(allowed), f"{perm} / {name}"


# ---------------------------------------------------------------- bibliothèque de templates
@pytest.mark.django_db
def test_template_library_crud_for_catalogue_managers(world):
    dom = Domaine.objects.first()
    body = {"domaine_id": dom.pk, "cas_usage": "Nouveau cas", "gravite": "Haute", "type_detection": "Simple",
            "exclusions": "e", "recommandation": "r", "impact": "i", "sources": [{"source_id": SourceLog.objects.first().pk}]}
    for u in (world.lecteur, world.adminclient, world.valideur):
        assert api(u).post("/api/templates", body, format="json").status_code == 403
    m = api(world.catalogue)
    r = m.post("/api/templates", body, format="json")
    assert r.status_code == 201
    code = r.json()["code"]
    assert code == "DET-083" and r.json()["impact"] == "i"  # code attribué automatiquement
    assert m.patch(f"/api/templates/{code}", {"impact": "modifié", "techniques": ["T1078"]}, format="json").json()["impact"] == "modifié"
    assert m.put(f"/api/templates/{code}/syntaxes/fortisiem", {"contenu": "<rule/>"}, format="json").status_code == 200
    dup = m.post(f"/api/templates/{code}/dupliquer", {}, format="json").json()
    assert dup["code"] == "DET-084" and len(dup["syntaxes"]) == 1
    assert m.delete(f"/api/templates/{code}").status_code == 204
    assert api(world.lecteur).get(f"/api/templates/{dup['code']}").json()["peut_modifier"] is False
    assert api(world.lecteur).patch(f"/api/templates/{dup['code']}", {"impact": "x"}, format="json").status_code == 403


@pytest.mark.django_db
def test_platform_referential_is_editable(world):
    m = api(world.catalogue)
    r = m.post("/api/siems", {"code": "monsiem", "nom": "Mon SIEM", "categorie": "siem", "langage_regle": "Maison", "langage_hunting": "Maison", "mode_editeur": "text", "ordre": 99}, format="json")
    assert r.status_code == 201
    assert m.delete(f"/api/siems/{r.json()['id']}").status_code == 204
    wazuh = next(s for s in m.get("/api/siems").json() if s["code"] == "wazuh")
    assert wazuh["nb_clients"] == 2
    assert api(world.lecteur).post("/api/siems", {"code": "x", "nom": "x"}, format="json").status_code == 403


@pytest.mark.django_db
def test_catalogue_stats_matches_excel_counters(world):
    s = api(world.lecteur).get("/api/catalogue/stats").json()
    assert s["total"] == 82
    row = next(d for d in s["domaines"] if d["domaine"] == "Identité & AD")
    assert (row["total"], row["Simple"], row["Seuil"], row["Corrélation"]) == (10, 3, 3, 4)  # = COUNTIFS du Lisez-moi
    assert sum(d["Simple"] for d in s["domaines"]) == 39
    assert sum(d["Seuil"] for d in s["domaines"]) == 18
    assert sum(d["Corrélation"] for d in s["domaines"]) == 25


@pytest.mark.django_db
def test_deleting_a_template_with_rules_and_hunting(world):
    """Régression : la suppression en cascade des règles ne doit pas échouer dans le journal d'audit."""
    m = api(world.catalogue)
    dom = Domaine.objects.first()
    code = m.post("/api/templates", {"domaine_id": dom.pk, "cas_usage": "À supprimer"}, format="json").json()["code"]
    assert m.put(f"/api/templates/{code}/syntaxes/fortisiem", {"contenu": "<rule/>"}, format="json").status_code == 200
    assert m.post(f"/api/templates/{code}/hunting", {"siem": "wazuh", "titre": "h", "requete": "q"}, format="json").status_code == 201
    assert m.get(f"/api/templates/{code}").status_code == 200  # la lecture précède la suppression, comme dans l'interface
    assert m.delete(f"/api/templates/{code}").status_code == 204
    assert not Template.objects.filter(code=code).exists()


@pytest.mark.django_db
def test_referentials_keep_their_configured_order(world):
    siems = [s["code"] for s in api(world.lecteur).get("/api/siems").json()]
    assert siems[:6] == ["wazuh", "splunk", "sentinel", "elastic", "qradar", "fortisiem"]
    domaines = [d["nom"] for d in api(world.lecteur).get("/api/domaines").json()]
    assert domaines[0] == "Identité & AD"


@pytest.mark.django_db
def test_seed_conseils_covers_every_template_and_never_overwrites(world):
    from django.core.management import call_command

    from catalog.content import CONSEILS

    assert set(CONSEILS) == set(Template.objects.values_list("code", flat=True))  # un texte pour chacun des 82 templates
    assert all(all(len(t[c].strip()) > 40 for c in ("exclusions", "recommandation", "impact")) for t in CONSEILS.values())
    det = get_det(world.a, "DET-002")
    api(world.contrib).patch(f"/api/detections/{det.pk}", {"impact": "Texte du client"}, format="json")
    call_command("seed_conseils")
    t = Template.objects.get(code="DET-002")
    assert t.exclusions and t.recommandation and t.impact
    det.refresh_from_db()
    assert det.impact == "Texte du client" and det.exclusions and det.recommandation  # le texte saisi est conservé, le reste est rempli
    t.exclusions = "Modifié"
    t.save()
    call_command("seed_conseils")
    t.refresh_from_db()
    assert t.exclusions == "Modifié"
